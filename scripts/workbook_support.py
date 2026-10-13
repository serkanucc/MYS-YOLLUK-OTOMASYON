from io import BytesIO
import re
import shutil
from zipfile import BadZipFile
from pathlib import Path

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.styles import Font, PatternFill


PERSONNEL_HEADERS = (
    "SIRANO",
    "TC NO",
    "AD",
    "SOYAD",
    "UNVAN",
    "GELDİĞİ KURUM",
    "GEÇ. GÖREV ÇALIŞTIĞI KURUM",
    "DERECE",
    "KADEME",
    "EK GÖSTERGE",
    "GÜNDELİĞİ",
    "GEÇ. GÖREV TARİHİ",
    "GÜN SAYISI",
    "IBAN",
    "Yolluk Tutar",
    "YOLLUK NO",
)


def validate_workbook(data):
    # Yüklenen Excel'in motorun beklediği iki sayfayı ve alan başlıklarını içerdiğini doğrula.
    try:
        workbook = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError) as exc:
        raise ValueError("Dosya geçerli veya okunabilir bir .xlsx çalışma kitabı değil.") from exc
    try:
        if len(workbook.worksheets) < 2:
            raise ValueError("Excel dosyasında Personel ve Birim Amiri olmak üzere en az 2 sayfa olmalı.")

        headers = {
            str(value).strip().upper()
            for value in next(workbook.worksheets[0].iter_rows(max_row=1, values_only=True), ())
            if value is not None
        }
        required_groups = (
            ("TC NO", "TCKN"),
            ("GEÇ. GÖREV TARİHİ", "GEÇ GÖREV TARİHİ"),
            ("GÜN SAYISI", "GUN SAYISI"),
            ("IBAN",),
            ("GÜNDELİĞİ", "GUNDELIGI"),
            ("SIRANO", "SIRA NO"),
        )
        missing = ["/".join(group) for group in required_groups if not headers.intersection(group)]
        if missing:
            raise ValueError("Personel sayfasında zorunlu sütun(lar) eksik: " + ", ".join(missing))

        amir_sheet = workbook.worksheets[1]
        amir_label_found = any(
            value is not None and re.search(r"(?:TC|TCKN)$", str(value).strip(), re.IGNORECASE)
            for row in amir_sheet.iter_rows(min_row=1, max_row=min(amir_sheet.max_row, 20), values_only=True)
            for value in row
        )
        if not amir_label_found:
            raise ValueError("Birim Amiri sayfasının ilk 20 satırında TC/TCKN etiketi bulunamadı.")
    finally:
        workbook.close()


def create_template():
    # Yolluk Liste 7 alanlarını koru; gerçek personel/kurum/amir verisini şablona kopyalama.
    workbook = openpyxl.Workbook()
    personnel = workbook.active
    personnel.title = "Personel"
    personnel.append(PERSONNEL_HEADERS)
    personnel.freeze_panes = "A2"
    personnel.auto_filter.ref = f"A1:{personnel.cell(row=1, column=len(PERSONNEL_HEADERS)).column_letter}1"
    personnel.row_dimensions[1].height = 30
    for cell in personnel[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for column, width in {
        "A": 12, "B": 16, "C": 18, "D": 18, "E": 22, "F": 30,
        "G": 36, "H": 12, "I": 12, "J": 16, "K": 16, "L": 24,
        "M": 14, "N": 30, "O": 16, "P": 24,
    }.items():
        personnel.column_dimensions[column].width = width

    manager = workbook.create_sheet("Birim Amiri")
    manager["A1"] = "Birim Amiri"
    manager["B1"] = "TC/TCKN"
    manager["C1"] = ""
    manager["B1"].font = Font(bold=True, color="FFFFFF")
    manager["B1"].fill = PatternFill("solid", fgColor="1F4E78")
    manager.column_dimensions["A"].width = 20
    manager.column_dimensions["B"].width = 16
    manager.column_dimensions["C"].width = 20

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def initialize_result_workbook(source, destination):
    # Sonuçları kaynak Excel'i değiştirmeyen ayrı bir kopyada biriktir.
    source = Path(source)
    destination = Path(destination)
    if source.resolve() == destination.resolve():
        raise ValueError("Sonuç Excel'i kaynak dosyanın üzerine yazılamaz.")
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, destination)


def write_yolluk_number(workbook_path, excel_row, yolluk_no):
    # Simüle edilen numarayı gerçek Excel satırına yaz; eksik sonuç sütununu sona ekle.
    from copy import copy

    workbook_path = Path(workbook_path)
    if excel_row < 2:
        raise ValueError("Yolluk numarası yalnızca personel veri satırlarına yazılabilir.")

    workbook = openpyxl.load_workbook(workbook_path)
    temporary = workbook_path.with_name(workbook_path.stem + ".tmp" + workbook_path.suffix)
    try:
        personnel = workbook[workbook.sheetnames[0]]
        headers = {
            str(cell.value).strip().upper(): cell.column
            for cell in personnel[1]
            if cell.value is not None
        }
        column = headers.get("YOLLUK NO")
        if column is None:
            column = personnel.max_column + 1
            header = personnel.cell(row=1, column=column)
            header.value = "YOLLUK NO"
            if column > 1:
                header._style = copy(personnel.cell(row=1, column=column - 1)._style)
            header.font = Font(bold=True, color="FFFFFF")
            header.fill = PatternFill("solid", fgColor="1F4E78")
            personnel.column_dimensions[header.column_letter].width = 24

        personnel.cell(row=excel_row, column=column).value = yolluk_no
        workbook.save(temporary)
        workbook.close()
        temporary.replace(workbook_path)
    except Exception:
        workbook.close()
        if temporary.exists():
            temporary.unlink()
        raise


def simulated_yolluk_number():
    # TEST- öneki, üretilen değerin MYS tarafından atanmış gerçek numara olmadığını belirtir.
    from uuid import uuid4
    from datetime import datetime

    return "TEST-" + datetime.now().strftime("%Y%m%d-%H%M%S") + "-" + uuid4().hex[:8].upper()
