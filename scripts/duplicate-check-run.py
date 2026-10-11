import json,os,re,shutil,subprocess,time
from pathlib import Path
import openpyxl
ROOT=Path(__file__).resolve().parents[1]
CFG=json.loads((ROOT/"config/runner-settings.json").read_text(encoding="utf-8"))
T=(ROOT/"scripts/duplicate-check.js").read_text(encoding="utf-8")


def resolve_cli():
    cfg=CFG.get("cliPath") if isinstance(CFG,dict) else None
    if cfg and Path(cfg).exists():
        return str(cfg)
    which=shutil.which("playwright-cli.cmd") or shutil.which("playwright-cli")
    if which:
        return which
    appdata=os.environ.get("APPDATA")
    if appdata:
        cand=Path(appdata)/"npm"/"playwright-cli.cmd"
        if cand.exists():
            return str(cand)
    return None


CLI=resolve_cli()
if not CLI:
    raise SystemExit("playwright-cli.cmd bulunamadı. Playwright CLI kurulumunu kontrol edin.")

_excel_cfg=CFG["sourceWorkbook"]
EXCEL=(ROOT/_excel_cfg) if not Path(_excel_cfg).is_absolute() else Path(_excel_cfg)
if not EXCEL.exists():
    raise SystemExit("Excel dosyası bulunamadı: "+str(EXCEL))
wb=openpyxl.load_workbook(EXCEL,data_only=True,read_only=True)
ws=wb[wb.sheetnames[0]]
rows=list(ws.iter_rows(values_only=True));h=[str(x).strip() if x is not None else "" for x in rows[0]]
i=h.index("TC NO") if "TC NO" in h else h.index("TCKN")
t=str(rows[1][i])
js=re.sub(r"\s+"," ",T.replace("{{TCKN}}",t))
subprocess.run([CLI,"-s="+str(CFG.get("browserSession","chrome-main")),"goto","https://butunlesik.hmb.gov.tr/hys/mys-yollukislemleri/yolluk"],check=True)
time.sleep(1)
p=subprocess.run([CLI,"-s="+str(CFG.get("browserSession","chrome-main")),"eval",js],capture_output=True,text=True,encoding="utf-8",errors="replace")
if p.returncode: raise SystemExit(p.stderr or p.stdout)
m=re.search(r"### Result\s*\n([\s\S]*?)(?:\n### Ran|$)",p.stdout)
print(m.group(1).strip() if m else p.stdout)
