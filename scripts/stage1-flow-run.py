import json,re,subprocess,time,datetime
from pathlib import Path
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
CFG=json.loads((ROOT/"config/runner-settings.json").read_text(encoding="utf-8"))
PW_HEALTH=ROOT.parent/"PLAYWRIGHT-CHROME-BAGLANTI"/"PW-GREEN-CHECK.ps1"


def playwright_health():
    try:
        green=subprocess.run(["powershell","-NoProfile","-ExecutionPolicy","Bypass","-File",str(PW_HEALTH)],capture_output=True,text=True,encoding="utf-8",errors="replace",timeout=10)
        return green.stdout.strip().splitlines()[-1:] == ["1"]
    except Exception:
        return False


def wait_for_playwright():
    announced=False
    while not playwright_health():
        if not announced:
            print("PLAYWRIGHT_WAIT | Yeşil Playwright bağlantısı bekleniyor...",flush=True)
            announced=True
        time.sleep(2)
    if announced: print("PLAYWRIGHT_RECONNECTED | Bağlantı yeniden kuruldu.",flush=True)

FIXED=json.loads((ROOT/"config/fixed-values.json").read_text(encoding="utf-8"))
TEMPLATE=(ROOT/"scripts/stage1-flow.js").read_text(encoding="utf-8")
EXCEL=Path(CFG["sourceWorkbook"])

def pw(*args):
    wait_for_playwright()
    p=subprocess.run([r"C:\Users\serkan.suc\AppData\Roaming\npm\playwright-cli.cmd","-s="+str(CFG.get("browserSession","chrome-main")),*args],capture_output=True,text=True,encoding="utf-8",errors="replace")
    if p.returncode: raise RuntimeError(p.stderr or p.stdout)
    return p.stdout

def ev(expr):
    out=pw("eval",expr)
    m=re.search(r"### Result\s*\n([\s\S]*?)(?:\n### Ran|$)",out)
    if not m: raise RuntimeError(out)
    try:return json.loads(m.group(1).strip())
    except:return m.group(1).strip()

def read_amir_tckn(wb):
    if len(wb.worksheets)<2: raise RuntimeError("Excel 2. sayfa yok: Birim Amiri TCKN bilgisi alınamadı.")
    ws=wb.worksheets[1]
    for row in ws.iter_rows(min_row=1,max_row=min(ws.max_row,20),values_only=True):
        for j,v in enumerate(row):
            label=str(v).strip() if v is not None else ""
            if label and re.search(r"(?:TC|TCKN)$",label,re.I) and j+1<len(row):
                val=row[j+1]; digits=re.sub(r"\\D","",str(val)) if val is not None else ""
                if len(digits)==11:return digits
    raise RuntimeError("Excel 2. sayfada Birim Amiri TC/TCKN bulunamadı. Güncel amir TCKN'sini 2. sayfaya girin.")

def people():
    wb=openpyxl.load_workbook(EXCEL,data_only=True,read_only=True);ws=wb[wb.sheetnames[0]]
    amir_tckn=read_amir_tckn(wb)
    rows=list(ws.iter_rows(values_only=True));h=[str(x).strip() if x is not None else "" for x in rows[0]]
    def col(*names):
        for n in names:
            if n in h:return h.index(n)
        raise RuntimeError("Excel sütunu eksik: "+names[0])
    ci,cd,ci2,cn,ct,cc=col("TC NO","TCKN"),col("GEÇ. GÖREV TARİHİ","GEÇ GÖREV TARİHİ"),col("IBAN"),col("SIRANO","SIRA NO"),col("GÜNDELİĞİ","GUNDELIGI"),None
    out=[]
    for r in rows[1:]:
        if not any(x is not None and str(x).strip() for x in r):continue
        raw=str(r[cd]).strip();m=re.fullmatch(r"(.+)/(\d{2})/(\d{4})",raw)
        days=[int(x.strip()) for x in m.group(1).split(",") if x.strip()]
        start=f"{min(days):02d}/{int(m.group(2)):02d}/{m.group(3)}"
        end=(datetime.date(int(m.group(3)),int(m.group(2)),max(days))+datetime.timedelta(days=1)).strftime("%d/%m/%Y")
        out.append({"row":r[cn] if cn is not None else len(out)+2,"tckn":r[ci],"iban":r[ci2],"daily":r[ct],"start":start,"end":end,"amirTckn":amir_tckn})
    return out

def main():
    p=people()
    pw("goto",CFG["formUrl"]);time.sleep(1)
    x=p[0]
    js=TEMPLATE.replace("{{TCKN}}",str(x["tckn"])).replace("{{IBAN}}",str(x["iban"])).replace("{{DAILY}}",str(x["daily"])).replace("{{START}}",x["start"]).replace("{{END}}",x["end"]).replace("{{AMIR_TCKN}}",str(x["amirTckn"])).replace("{{BUDGET}}",FIXED["budgetItem"])
    js=re.sub(r"\s+"," ",js)
    tmp=ROOT/"data/last-flow.js"
    tmp.write_text(js,encoding="utf-8")
    chk=subprocess.run(["node","--check",str(tmp)],capture_output=True,text=True,encoding="utf-8",errors="replace")
    if chk.returncode: raise RuntimeError(chk.stderr)
    result=ev(js)
    print(json.dumps({"row":x["row"],"safeDryRun":True,"result":result},ensure_ascii=False))
if __name__=="__main__":main()
