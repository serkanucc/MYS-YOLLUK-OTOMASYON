"""Test amaçlı SAHTE personel Excel'i üretir: data\\Yolluk Liste 7.xlsx

UYARI: Bu dosya gerçek kişi verisi İÇERMEZ. IBAN'lar MOD-97 açısından geçerli
hesaplanır ama bankada kayıtlı değildir; TCKN'ler sahtedir. Amaç yalnızca
Excel okuma / IBAN ön kontrol / tarih çözümleme / assignmentId mantığını
MYS'e gitmeden yerelde doğrulamaktır.

Kullanım:  python scripts\\make_test_excel.py
Çıktı    : data\\Yolluk Liste 7.xlsx  (.gitignore gereği git'e eklenmez)
"""
import random
import re
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from assignment import duty_period  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "Yolluk Liste 7.xlsx"

HEADERS = ["SIRANO", "TC NO", "AD", "SOYAD", "UNVAN", "DERECE", "KADEME",
           "EK GÖSTERGE", "GÜNDELİĞİ", "GEÇ. GÖREV TARİHİ", "GÜN SAYISI", "IBAN"]

# assignment.py'nin desteklediği farklı görev tarihi biçimleri
DUTY_FORMATS = [
    "03,10,12/08/2026",       # virgüllü gün listesi
    "01,05,08/09/2026",       # virgüllü gün listesi
    "10-20/10/2026",          # gün aralığı
    "01.11.2026,05.11.2026",  # tam tarih listesi (nokta ayraç)
    "15-18/12/2026",          # gün aralığı
    "22,25/01/2027",          # farklı yıl
]

BANKS = ["00061", "00032", "00046", "00059", "00012", "00010", "00015", "00006"]
TITLES = ["Uzman Çavuş", "Hemşire", "Ebe", "Radyoloji Teknikeri", "Sağlık Teknikeri"]


def _mod97_convert(s):
    out = ""
    for ch in s:
        out += ch if ch.isdigit() else str(ord(ch) - 55)
    return int(out)


def make_tr_iban(bank):
    """MOD-97 geçerli (ama sahte) TR IBAN üretir."""
    acct = bank + "".join(random.choice("0123456789") for _ in range(17))
    check = 98 - (_mod97_convert(acct + "TR00") % 97)
    return "TR" + f"{check:02d}" + acct


def valid_tr_iban(value):
    """stage1-live-run.py içindeki kontrolün aynısı."""
    s = re.sub(r"\s+", "", str(value or "").upper())
    return bool(re.fullmatch(r"TR\d{24}", s)) and _mod97_convert(s[4:] + s[:4]) % 97 == 1


def main():
    random.seed(20261011)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Personel"
    ws.append(HEADERS)

    for i, duty in enumerate(DUTY_FORMATS, 1):
        bank = BANKS[(i - 1) % len(BANKS)]
        iban = make_tr_iban(bank)
        tckn = str(random.randint(1, 9)) + "".join(random.choice("0123456789") for _ in range(10))
        dates, _start, _end = duty_period(duty)  # tarih çözümlemesi burada doğrulanır
        ws.append([
            i,
            tckn,
            f"Test{i}",
            f"Kisi{i}",
            TITLES[(i - 1) % len(TITLES)],
            random.randint(1, 15),
            random.randint(1, 4),
            random.randint(100, 3000),
            850.00,
            duty,
            len(dates),
            iban,
        ])

    # 2. sayfa: Birim Amiri TCKN (TC/TCKN etiketinin sağındaki 11 haneli değer)
    ws2 = wb.create_sheet("Birim Amiri")
    ws2.append(["Birim Amiri TCKN", "12345678901"])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)

    # Geri okuyup kendi ürettiği veriyi doğrula
    chk = openpyxl.load_workbook(OUT, data_only=True, read_only=True)
    ws_v = chk[chk.sheetnames[0]]
    rows = list(ws_v.iter_rows(values_only=True))
    assert rows[0] == tuple(HEADERS), "Başlıklar hatalı"
    n = 0
    for r in rows[1:]:
        assert valid_tr_iban(r[11]), f"Geçersiz IBAN üretildi: {r[11]}"
        _d, _s, _e = duty_period(r[9])
        n += 1
    assert len(chk.sheetnames) >= 2, "2. sayfa (Birim Amiri) yok"

    print(f"Yazıldı: {OUT}")
    print(f"Personel satırı: {n}  |  Birim Amiri sayfası: {chk.sheetnames[1]}")
    print("Tüm IBAN'lar MOD-97 geçerli, tüm tarih biçimleri çözüldü. TEST VERİSİ HAZIR.")


if __name__ == "__main__":
    main()
