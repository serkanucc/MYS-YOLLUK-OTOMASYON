import sys
import importlib.util
import importlib.util
import sys
from io import BytesIO
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from workbook_support import (
    create_template,
    initialize_result_workbook,
    simulated_yolluk_number,
    validate_workbook,
    write_yolluk_number,
)

# Şablonun ikinci sayfası motorun mevcut amir TCKN okuyucusuyla da uyumlu olmalı.
runner_path = Path(__file__).resolve().parents[1] / "scripts" / "stage1-live-run.py"
spec = importlib.util.spec_from_file_location("stage1_live_run", runner_path)
runner = importlib.util.module_from_spec(spec)
spec.loader.exec_module(runner)

# İndirilen boş şablon upload kontrolünden ve motorun amir etiketi kuralından geçmeli.
template = create_template()
validate_workbook(template)
template_workbook = openpyxl.load_workbook(BytesIO(template))
template_headers = [cell.value for cell in template_workbook.worksheets[0][1]]
assert template_headers == [
    "SIRANO", "TC NO", "AD", "SOYAD", "UNVAN", "GELDİĞİ KURUM",
    "GEÇ. GÖREV ÇALIŞTIĞI KURUM", "DERECE", "KADEME", "EK GÖSTERGE",
    "GÜNDELİĞİ", "GEÇ. GÖREV TARİHİ", "GÜN SAYISI", "IBAN",
    "Yolluk Tutar", "YOLLUK NO",
]
assert template_workbook.worksheets[0].max_row == 1
template_workbook.close()

# Gün sayısı doğrulaması için başlık Excel yüklemesinde zorunlu olmalı.
without_day_count = openpyxl.load_workbook(BytesIO(template))
without_day_count.worksheets[0].delete_cols(13)
buffer = BytesIO()
without_day_count.save(buffer)
without_day_count.close()
try:
    validate_workbook(buffer.getvalue())
except ValueError as error:
    assert "GÜN SAYISI" in str(error)
else:
    raise AssertionError("GÜN SAYISI başlığı olmayan Excel kabul edildi")

try:
    runner.read_amir_tckn(openpyxl.load_workbook(BytesIO(template)))
except RuntimeError as error:
    assert "Birim Amiri TC/TCKN bulunamadı" in str(error)
else:
    raise AssertionError("Boş şablonda amir TCKN doldurulmadığı halde hata üretilmedi")

# Yanlış sütunlu Excel, yükleme sırasında açıklayıcı bir hatayla reddedilmeli.
invalid = openpyxl.Workbook()
invalid.create_sheet("Birim Amiri")
buffer = BytesIO()
invalid.save(buffer)
try:
    validate_workbook(buffer.getvalue())
except ValueError as error:
    assert "zorunlu sütun" in str(error)
else:
    raise AssertionError("Zorunlu sütunları olmayan çalışma kitabı kabul edildi")

# Test numarası yalnızca kaynak Excel'in kopyasına, hedef personel satırında yazılmalı.
source = openpyxl.Workbook()
personnel = source.active
personnel.title = "Personel"
personnel.append(["SIRANO", "TC NO"])
personnel.append([10, "TEST-TCKN"])
source.create_sheet("Birim Amiri")
source_path = Path(__file__).resolve().parent / "_test-source.xlsx"
result_path = Path(__file__).resolve().parent / "_test-result.xlsx"
try:
    source.save(source_path)
    source.close()
    initialize_result_workbook(source_path, result_path)
    number = simulated_yolluk_number()
    assert number.startswith("TEST-")
    write_yolluk_number(result_path, 2, number)
    result = openpyxl.load_workbook(result_path, data_only=True)
    assert result["Personel"].cell(1, 3).value == "YOLLUK NO"
    assert result["Personel"].cell(2, 3).value == number
    result.close()
    original = openpyxl.load_workbook(source_path, data_only=True)
    assert original["Personel"].cell(2, 3).value is None
    original.close()
finally:
    source_path.unlink(missing_ok=True)
    result_path.unlink(missing_ok=True)

print("WORKBOOK SUPPORT TESTS OK")
